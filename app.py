import os
import re
import secrets
import shutil
import sqlite3
import tempfile
import uuid
import zipfile
from io import BytesIO
from datetime import datetime
from pathlib import Path

import pytesseract
from flask_wtf.csrf import CSRFProtect
from flask import Flask, flash, g, redirect, render_template, request, send_file, send_from_directory, session, url_for
from werkzeug.datastructures import FileStorage
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

# Tesseract detection - works on both Windows and Linux (Railway)
pytesseract.pytesseract.tesseract_cmd = shutil.which("tesseract") or r"C:\Program Files\Tesseract-OCR\tesseract.exe"

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config["DATABASE"] = os.environ.get("DATABASE_URL", os.path.join(app.root_path, "archive.db"))
app.config["UPLOAD_FOLDER"] = os.environ.get("UPLOAD_FOLDER", os.path.join(os.path.dirname(__file__), "uploads"))
app.config["TESSDATA_DIR"] = os.environ.get("TESSDATA_DIR") or os.path.join(app.root_path, "tessdata")
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
csrf = CSRFProtect(app)

os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)

CATEGORY_KEYWORDS = {
    "Surat Masuk": ["surat masuk", "diterima", "dari", "masuk"],
    "Surat Keluar": ["surat keluar", "dikirim", "kepada", "keluar"],
    "Surat Undangan": ["undangan", "rapat", "pelaksanaan", "acara"],
    "Surat Keterangan": ["keterangan", "domisili", "usaha", "tidak mampu", "kelahiran"],
    "Surat Pemberitahuan": ["pemberitahuan", "pengumuman", "informasi", "berita"],
    "Surat Administrasi Lainnya": ["administrasi", "permohonan", "nota dinas", "laporan", "persetujuan"],
}


def get_db():
    if "db" not in g:
        db_url = app.config["DATABASE"]
        if db_url.startswith("postgresql://") or db_url.startswith("postgres://"):
            try:
                import psycopg2
                from psycopg2.extras import RealDictCursor
                conn = psycopg2.connect(db_url)
                conn.row_factory = RealDictCursor
                g.db = conn
                return g.db
            except ImportError:
                app.logger.warning("psycopg2 not available, falling back to SQLite")
                db_url = os.path.join(app.root_path, "archive.db")
                app.config["DATABASE"] = db_url
        conn = sqlite3.connect(db_url)
        conn.row_factory = sqlite3.Row
        g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(_error):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def log_activity(db, action, details="", archive_id=None):
    db.execute(
        "INSERT INTO activity_logs (username, action, details, archive_id) VALUES (?, ?, ?, ?)",
        (session.get("user", "system"), action, details, archive_id),
    )


def is_supported_upload(file_storage):
    filename = secure_filename(file_storage.filename or "")
    extension = os.path.splitext(filename)[1].lower()
    signatures = {
        ".pdf": lambda header: header.startswith(b"%PDF-"),
        ".png": lambda header: header.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": lambda header: header.startswith(b"\xff\xd8\xff"),
        ".jpeg": lambda header: header.startswith(b"\xff\xd8\xff"),
    }
    signature_check = signatures.get(extension)
    if signature_check is None:
        return False

    file_storage.seek(0)
    header = file_storage.read(16)
    file_storage.seek(0)
    return signature_check(header)


def init_db():
    db = get_db()
    db_url = app.config["DATABASE"]
    is_postgres = db_url.startswith("postgresql://") or db_url.startswith("postgres://")
    
    if is_postgres:
        # PostgreSQL syntax
        db.execute("""
            CREATE TABLE IF NOT EXISTS archives (
                id SERIAL PRIMARY KEY,
                nomor_surat TEXT,
                tanggal TEXT,
                perihal TEXT,
                pengirim TEXT,
                tujuan TEXT,
                kategori TEXT,
                file_name TEXT,
                file_path TEXT,
                extracted_text TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                deleted_at TIMESTAMP
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL,
                is_active INTEGER NOT NULL DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS activity_logs (
                id SERIAL PRIMARY KEY,
                username TEXT NOT NULL,
                action TEXT NOT NULL,
                details TEXT,
                archive_id INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
    else:
        # SQLite syntax
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS archives (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nomor_surat TEXT,
                tanggal TEXT,
                perihal TEXT,
                pengirim TEXT,
                tujuan TEXT,
                kategori TEXT,
                file_name TEXT,
                file_path TEXT,
                extracted_text TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                deleted_at TEXT
            )
            """
        )
        archive_columns = {row["name"] for row in db.execute("PRAGMA table_info(archives)").fetchall()}
        if "deleted_at" not in archive_columns:
            db.execute("ALTER TABLE archives ADD COLUMN deleted_at TEXT")
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL,
                is_active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        user_columns = {row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()}
        if "is_active" not in user_columns:
            db.execute("ALTER TABLE users ADD COLUMN is_active INTEGER NOT NULL DEFAULT 1")
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS activity_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL,
                action TEXT NOT NULL,
                details TEXT,
                archive_id INTEGER,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
    
    default_users = (
        ("admin", os.environ.get("ARCHIVE_ADMIN_PASSWORD", "admin123"), "admin"),
        ("petugas", os.environ.get("ARCHIVE_STAFF_PASSWORD", "petugas123"), "petugas"),
    )
    for username, password, role in default_users:
        existing_user = db.execute(
            "SELECT 1 FROM users WHERE username = ?", (username,)
        ).fetchone()
        if existing_user is None:
            db.execute(
                "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                (username, generate_password_hash(password), role),
            )
    db.commit()


with app.app_context():
    init_db()


@app.before_request
def validate_user_session():
    username = session.get("user")
    if not username:
        return None

    user = get_db().execute(
        "SELECT role, is_active FROM users WHERE username = ?", (username,)
    ).fetchone()
    if user is None or not user["is_active"]:
        session.clear()
        flash("Akun tidak aktif atau tidak ditemukan. Silakan login kembali.", "warning")
        if request.endpoint not in ("login", "static", "favicon"):
            return redirect(url_for("login"))
        return None

    session["role"] = user["role"]
    return None


def normalize_text(text):
    return re.sub(r"\s+", " ", text or "").strip()


def get_date_filters():
    filters = {}
    for name in ("start_date", "end_date"):
        value = request.args.get(name, "").strip()
        try:
            datetime.strptime(value, "%Y-%m-%d")
            filters[name] = value
        except ValueError:
            filters[name] = ""
    return filters


def archive_filter_query(query="", category="", start_date="", end_date=""):
    conditions = ["deleted_at IS NULL"]
    params = []
    if query:
        conditions.append(
            "(nomor_surat LIKE ? OR perihal LIKE ? OR kategori LIKE ? OR pengirim LIKE ? OR tujuan LIKE ? OR extracted_text LIKE ?)"
        )
        params.extend([f"%{query}%"] * 6)
    if category:
        conditions.append("kategori = ?")
        params.append(category)
    if start_date:
        conditions.append("tanggal >= ?")
        params.append(start_date)
    if end_date:
        conditions.append("tanggal <= ?")
        params.append(end_date)

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    return where_clause, params


def parse_date(value):
    value = (value or "").strip()
    if not value:
        return ""
    value = re.sub(
        r"^(?:senin|selasa|rabu|kamis|jumat|jum'at|sabtu|minggu)\s*,?\s*",
        "",
        value,
        flags=re.IGNORECASE,
    )
    date_match = re.search(
        r"\b\d{1,2}(?:\s+|[-/])(?:[A-Za-z]+|\d{1,2})(?:\s*,?\s*|[-/])\d{4}\b",
        value,
    )
    if date_match:
        value = date_match.group(0).strip()

    formats = [
        "%d %B %Y",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d %b %Y",
        "%d %B, %Y",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    months = {
        "januari": "January",
        "februari": "February",
        "maret": "March",
        "april": "April",
        "mei": "May",
        "juni": "June",
        "juli": "July",
        "agustus": "August",
        "september": "September",
        "oktober": "October",
        "november": "November",
        "desember": "December",
    }
    translated = re.sub(
        r"\b(" + "|".join(months) + r")\b",
        lambda match: months[match.group(0).lower()],
        value,
        flags=re.IGNORECASE,
    )
    for fmt in ("%d %B %Y", "%d %B, %Y"):
        try:
            return datetime.strptime(translated, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return value


def normalize_archive_dates():
    db = get_db()
    records = db.execute(
        "SELECT id, tanggal FROM archives WHERE tanggal IS NOT NULL AND tanggal != ''"
    ).fetchall()
    updates = []
    for record in records:
        normalized = parse_date(record["tanggal"])
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", normalized) and normalized != record["tanggal"]:
            updates.append((normalized, record["id"]))
    if updates:
        db.executemany("UPDATE archives SET tanggal = ? WHERE id = ?", updates)
        db.commit()


with app.app_context():
    normalize_archive_dates()


def extract_metadata(text):
    cleaned = text or ""
    metadata = {
        "nomor_surat": "",
        "tanggal": "",
        "perihal": "",
        "pengirim": "",
        "tujuan": "",
    }

    patterns = {
        "nomor_surat": [
            r"nomor\s*surat\s*[:\-]?\s*([^\n]+)",
            r"no\.?\s*surat\s*[:\-]?\s*([^\n]+)",
            r"no\.?\s*[:\-]?\s*([^\n]+)",
        ],
        "tanggal": [
            r"tanggal\s*[:\-]?\s*([^\n]+)",
            r"tgl\.?\s*[:\-]?\s*([^\n]+)",
        ],
        "perihal": [
            r"perihal\s*[:\-]?\s*([^\n]+)",
            r"hal\s*[:\-]?\s*([^\n]+)",
        ],
        "pengirim": [
            r"pengirim\s*[:\-]?\s*([^\n]+)",
            r"dari\s*[:\-]?\s*([^\n]+)",
        ],
        "tujuan": [
            r"tujuan\s*[:\-]?\s*([^\n]+)",
            r"kepada\s*[:\-]?\s*([^\n]+)",
        ],
    }

    for key, regexes in patterns.items():
        for pattern in regexes:
            match = re.search(pattern, cleaned, flags=re.IGNORECASE)
            if match:
                value = match.group(1).strip(" .:-")
                metadata[key] = value
                break

    if metadata["tanggal"]:
        metadata["tanggal"] = parse_date(metadata["tanggal"])

    if not metadata["nomor_surat"]:
        match = re.search(r"\b\d{1,4}/[A-Za-z0-9./-]+\b", cleaned)
        if match:
            metadata["nomor_surat"] = match.group(0)

    if not metadata["perihal"]:
        match = re.search(r"perihal\s*[:\-]?\s*(.+)", cleaned, flags=re.IGNORECASE)
        if match:
            metadata["perihal"] = match.group(1).strip()

    return metadata


def classify_document(text):
    lowercase = (text or "").lower()
    for category, keywords in CATEGORY_KEYWORDS.items():
        if any(keyword in lowercase for keyword in keywords):
            return category
    if "undangan" in lowercase:
        return "Surat Undangan"
    if "keterangan" in lowercase:
        return "Surat Keterangan"
    return "Surat Administrasi Lainnya"


def get_tesseract_config(options=""):
    tessdata_dir = app.config["TESSDATA_DIR"]
    # Check multiple possible locations for traineddata files
    possible_dirs = [
        tessdata_dir,
        "/usr/share/tesseract-ocr/5/tessdata",
        "/usr/share/tesseract-ocr/tessdata",
        "/usr/share/tesseract/5/tessdata",
        "/usr/share/tesseract/tessdata",
    ]
    for candidate in possible_dirs:
        if os.path.isdir(candidate):
            has_eng = os.path.isfile(os.path.join(candidate, "eng.traineddata"))
            has_ind = os.path.isfile(os.path.join(candidate, "ind.traineddata"))
            if has_eng and has_ind:
                return f'{options} --tessdata-dir "{candidate}"'.strip()
            elif has_eng or has_ind:
                return f'{options} --tessdata-dir "{candidate}"'.strip()
    return options


def get_ocr_languages():
    try:
        installed_languages = set(pytesseract.get_languages(config=get_tesseract_config()))
    except pytesseract.TesseractNotFoundError:
        installed_languages = set()
    if "ind" in installed_languages and "eng" in installed_languages:
        return "ind+eng"
    if "ind" in installed_languages:
        return "ind"
    return "eng"


def ocr_image(image):
    from PIL import ImageEnhance, ImageOps, ImageFilter

    image = ImageOps.exif_transpose(image).convert("L")
    image = image.resize((image.width * 3, image.height * 3))
    image = ImageOps.autocontrast(image)
    image = ImageEnhance.Contrast(image).enhance(1.5)
    image = image.filter(ImageFilter.SHARPEN)
    return pytesseract.image_to_string(
        image, lang=get_ocr_languages(), config=get_tesseract_config("--oem 3 --psm 6")
    )


def read_text_from_upload(file_storage):
    filename = (file_storage.filename or "").lower()
    text = ""

    if filename.endswith(".txt"):
        try:
            text = file_storage.read().decode("utf-8", errors="ignore")
            file_storage.seek(0)
        except Exception:
            pass

    if filename.endswith(".pdf"):
        try:
            import os
            import tempfile

            from pdf2image import convert_from_bytes
            import pytesseract

            file_storage.seek(0)
            pdf_bytes = file_storage.read()
            images = convert_from_bytes(pdf_bytes)
            page_texts = []
            for image in images:
                page_texts.append(
                    pytesseract.image_to_string(
                        image,
                        lang=get_ocr_languages(),
                        config=get_tesseract_config(),
                    )
                )
            text = "\n".join(page_texts)
            file_storage.seek(0)
        except Exception:
            try:
                import pypdf

                file_storage.seek(0)
                pdf = pypdf.PdfReader(file_storage)
                pages = [page.extract_text() or "" for page in pdf.pages]
                text = "\n".join(pages)
                file_storage.seek(0)
            except Exception:
                text = "Hasil OCR diproses secara otomatis untuk dokumen PDF."

    if filename.endswith((".png", ".jpg", ".jpeg")):
        try:
            from PIL import Image

            image = Image.open(file_storage)
            text = ocr_image(image)
            file_storage.seek(0)
        except Exception:
            app.logger.exception("Image OCR failed for %s", file_storage.filename)
            text = "OCR gambar gagal diproses. Periksa instalasi Tesseract dan kualitas gambar."

    return text or "Dokumen berhasil diunggah, namun teks belum otomatis terbaca."


@app.route("/favicon.ico")
def favicon():
    return send_from_directory(os.path.join(app.root_path, "static"), "favicon.ico", mimetype="image/x-icon")


@app.route("/")
def index():
    if "user" in session:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = get_db().execute(
            "SELECT username, password_hash, role, is_active FROM users WHERE username = ?",
            (username,),
        ).fetchone()

        if user and user["is_active"] and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user"] = username
            session["role"] = user["role"]
            flash("Login berhasil.", "success")
            return redirect(url_for("dashboard"))

        flash("Username atau password salah.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Anda telah logout.", "info")
    return redirect(url_for("login"))


@app.route("/change-password", methods=["GET", "POST"])
def change_password():
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    user = db.execute(
        "SELECT password_hash FROM users WHERE username = ?", (session["user"],)
    ).fetchone()
    if user is None:
        session.clear()
        flash("Akun tidak ditemukan. Silakan login kembali.", "danger")
        return redirect(url_for("login"))

    if request.method == "POST":
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not check_password_hash(user["password_hash"], current_password):
            flash("Password saat ini salah.", "danger")
        elif len(new_password) < 12:
            flash("Password baru harus terdiri dari minimal 12 karakter.", "danger")
        elif new_password != confirm_password:
            flash("Konfirmasi password baru tidak sama.", "danger")
        else:
            db.execute(
                "UPDATE users SET password_hash = ? WHERE username = ?",
                (generate_password_hash(new_password), session["user"]),
            )
            db.commit()
            session.clear()
            flash("Password berhasil diubah. Silakan login kembali.", "success")
            return redirect(url_for("login"))

    return render_template("change_password.html")


@app.route("/admin/users", methods=["GET", "POST"])
def manage_users():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat mengelola akun.", "danger")
        return redirect(url_for("dashboard"))

    db = get_db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "petugas")

        if not re.fullmatch(r"[A-Za-z0-9_.-]{3,50}", username):
            flash("Username harus 3-50 karakter dan hanya boleh berisi huruf, angka, titik, garis bawah, atau tanda hubung.", "danger")
        elif len(password) < 12:
            flash("Password awal harus terdiri dari minimal 12 karakter.", "danger")
        elif role not in ("admin", "petugas"):
            flash("Peran akun tidak valid.", "danger")
        elif db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
            flash("Username sudah digunakan.", "danger")
        else:
            db.execute(
                "INSERT INTO users (username, password_hash, role, is_active) VALUES (?, ?, ?, 1)",
                (username, generate_password_hash(password), role),
            )
            log_activity(db, "user_create", f"Membuat akun {username} ({role})")
            db.commit()
            flash("Akun berhasil dibuat.", "success")
            return redirect(url_for("manage_users"))

    users = db.execute(
        "SELECT username, role, is_active, created_at FROM users ORDER BY username"
    ).fetchall()
    return render_template("users.html", users=users)


@app.route("/admin/users/<username>/update", methods=["POST"])
def update_user(username):
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat mengelola akun.", "danger")
        return redirect(url_for("dashboard"))

    db = get_db()
    user = db.execute(
        "SELECT username, role, is_active FROM users WHERE username = ?", (username,)
    ).fetchone()
    if user is None:
        flash("Akun tidak ditemukan.", "danger")
        return redirect(url_for("manage_users"))

    role = request.form.get("role", "")
    is_active = 1 if request.form.get("is_active") == "1" else 0
    if role not in ("admin", "petugas"):
        flash("Peran akun tidak valid.", "danger")
        return redirect(url_for("manage_users"))
    if username == session["user"] and (role != "admin" or not is_active):
        flash("Admin tidak dapat menurunkan peran atau menonaktifkan akunnya sendiri.", "danger")
        return redirect(url_for("manage_users"))
    if user["role"] == "admin" and user["is_active"] and (role != "admin" or not is_active):
        active_admins = db.execute(
            "SELECT COUNT(*) FROM users WHERE role = 'admin' AND is_active = 1"
        ).fetchone()[0]
        if active_admins <= 1:
            flash("Setidaknya harus ada satu admin aktif.", "danger")
            return redirect(url_for("manage_users"))

    db.execute(
        "UPDATE users SET role = ?, is_active = ? WHERE username = ?",
        (role, is_active, username),
    )
    log_activity(db, "user_update", f"Memperbarui akun {username}: {role}, aktif={is_active}")
    db.commit()
    flash("Akun berhasil diperbarui.", "success")
    return redirect(url_for("manage_users"))


@app.route("/admin/users/<username>/password", methods=["POST"])
def reset_user_password(username):
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat mengelola akun.", "danger")
        return redirect(url_for("dashboard"))

    password = request.form.get("password", "")
    if len(password) < 12:
        flash("Password baru harus terdiri dari minimal 12 karakter.", "danger")
        return redirect(url_for("manage_users"))

    db = get_db()
    user = db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone()
    if user is None:
        flash("Akun tidak ditemukan.", "danger")
        return redirect(url_for("manage_users"))

    db.execute(
        "UPDATE users SET password_hash = ? WHERE username = ?",
        (generate_password_hash(password), username),
    )
    log_activity(db, "user_password_reset", f"Reset password akun {username}")
    db.commit()
    flash("Password akun berhasil direset.", "success")
    return redirect(url_for("manage_users"))


@app.route("/dashboard")
def dashboard():
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    total = db.execute(
        "SELECT COUNT(*) AS total FROM archives WHERE deleted_at IS NULL"
    ).fetchone()["total"]
    categories = db.execute(
        "SELECT kategori, COUNT(*) AS total FROM archives WHERE deleted_at IS NULL GROUP BY kategori ORDER BY total DESC"
    ).fetchall()
    recent = db.execute(
        "SELECT * FROM archives WHERE deleted_at IS NULL ORDER BY created_at DESC LIMIT 5"
    ).fetchall()
    return render_template("dashboard.html", total=total, categories=categories, recent=recent)


@app.route("/upload", methods=["GET", "POST"])
def upload_document():
    if "user" not in session:
        return redirect(url_for("login"))

    if request.method == "POST":
        file = request.files.get("document")
        if not file or file.filename == "":
            flash("Pilih dokumen yang akan diunggah.", "danger")
            return redirect(url_for("upload_document"))

        original_filename = secure_filename(file.filename)
        if not original_filename or not is_supported_upload(file):
            flash("Format file tidak didukung atau isi file tidak sesuai. Gunakan PDF, PNG, JPG, atau JPEG yang valid.", "danger")
            return redirect(url_for("upload_document"))

        extension = os.path.splitext(original_filename)[1].lower()
        stored_filename = f"{uuid.uuid4().hex}{extension}"
        saved_path = os.path.join(app.config["UPLOAD_FOLDER"], stored_filename)
        file.save(saved_path)

        file.seek(0)
        extracted_text = read_text_from_upload(file)
        metadata = extract_metadata(extracted_text)
        category = classify_document(extracted_text)
        selected_category = request.form.get("kategori") or category

        if request.form.get("nomor_surat"):
            metadata["nomor_surat"] = request.form.get("nomor_surat")
        if request.form.get("tanggal"):
            metadata["tanggal"] = parse_date(request.form.get("tanggal"))
        if request.form.get("perihal"):
            metadata["perihal"] = request.form.get("perihal")
        if request.form.get("pengirim"):
            metadata["pengirim"] = request.form.get("pengirim")
        if request.form.get("tujuan"):
            metadata["tujuan"] = request.form.get("tujuan")

        db = get_db()
        cursor = db.execute(
            """
            INSERT INTO archives (
                nomor_surat, tanggal, perihal, pengirim, tujuan, kategori,
                file_name, file_path, extracted_text
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                metadata.get("nomor_surat", ""),
                metadata.get("tanggal", ""),
                metadata.get("perihal", ""),
                metadata.get("pengirim", ""),
                metadata.get("tujuan", ""),
                selected_category,
                original_filename,
                saved_path,
                extracted_text,
            ),
        )
        log_activity(db, "upload", f"Unggah {original_filename}", cursor.lastrowid)
        db.commit()
        flash("Dokumen berhasil diunggah dan diproses.", "success")
        return redirect(url_for("archives"))

    return render_template("upload.html")


@app.route("/archives")
def archives():
    if "user" not in session:
        return redirect(url_for("login"))

    query = request.args.get("q", "")
    category = request.args.get("category", "")
    date_filters = get_date_filters()
    db = get_db()
    where_clause, params = archive_filter_query(
        query, category, date_filters["start_date"], date_filters["end_date"]
    )
    rows = db.execute(
        f"SELECT * FROM archives {where_clause} ORDER BY created_at DESC, id DESC", params
    ).fetchall()
    categories = db.execute(
        "SELECT DISTINCT kategori FROM archives WHERE deleted_at IS NULL AND kategori IS NOT NULL AND kategori != '' ORDER BY kategori"
    ).fetchall()
    return render_template(
        "archives.html",
        records=rows,
        query=query,
        categories=categories,
        selected_category=category,
        **date_filters,
    )


@app.route("/archives/reset", methods=["POST"])
def archives_reset():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat menghapus seluruh arsip.", "danger")
        return redirect(url_for("archives"))

    db = get_db()
    rows = db.execute("SELECT COUNT(*) FROM archives WHERE deleted_at IS NULL").fetchone()[0]
    db.execute("UPDATE archives SET deleted_at = CURRENT_TIMESTAMP WHERE deleted_at IS NULL")
    log_activity(db, "reset", f"Pindahkan seluruh arsip ke tempat sampah ({rows} dokumen)")
    db.commit()
    flash(f"{rows} arsip dipindahkan ke tempat sampah dan dapat dipulihkan.", "success")
    return redirect(url_for("archives"))


@app.route("/admin/activity")
def activity_log():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat melihat riwayat aktivitas.", "danger")
        return redirect(url_for("dashboard"))

    logs = get_db().execute(
        "SELECT * FROM activity_logs ORDER BY created_at DESC, id DESC LIMIT 500"
    ).fetchall()
    return render_template("activity_log.html", logs=logs)


@app.route("/admin/backup")
def download_backup():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat mengunduh backup.", "danger")
        return redirect(url_for("dashboard"))

    db = get_db()
    log_activity(db, "backup", "Membuat backup database dan dokumen arsip")
    db.commit()

    backup_buffer = BytesIO()
    backup_file = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    backup_path = backup_file.name
    backup_file.close()
    try:
        backup_db = sqlite3.connect(backup_path)
        try:
            db.backup(backup_db)
        finally:
            backup_db.close()

        upload_folder = os.path.realpath(app.config["UPLOAD_FOLDER"])
        attachments = db.execute(
            "SELECT DISTINCT file_path FROM archives WHERE file_path IS NOT NULL AND file_path != ''"
        ).fetchall()
        with zipfile.ZipFile(backup_buffer, "w", zipfile.ZIP_DEFLATED) as archive_zip:
            archive_zip.write(backup_path, "archive.db")
            for attachment in attachments:
                file_path = os.path.realpath(attachment["file_path"])
                try:
                    is_uploaded_file = os.path.commonpath((upload_folder, file_path)) == upload_folder
                except ValueError:
                    is_uploaded_file = False
                if is_uploaded_file and os.path.isfile(file_path):
                    relative_path = os.path.relpath(file_path, upload_folder).replace(os.sep, "/")
                    archive_zip.write(file_path, f"uploads/{relative_path}")
    finally:
        if os.path.exists(backup_path):
            os.remove(backup_path)

    backup_buffer.seek(0)
    return send_file(
        backup_buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name=f"backup_arsip_{datetime.now():%Y%m%d_%H%M%S}.zip",
    )


@app.route("/admin/backup/restore", methods=["GET", "POST"])
def restore_backup():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat memulihkan backup.", "danger")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        file = request.files.get("backup_file")
        if not file or file.filename == "":
            flash("Pilih file backup yang akan dipulihkan.", "danger")
            return redirect(url_for("restore_backup"))

        try:
            with tempfile.TemporaryDirectory() as tmp_dir:
                archive_path = os.path.join(tmp_dir, "archive.db")
                with zipfile.ZipFile(file, "r") as backup_zip:
                    if "archive.db" not in backup_zip.namelist():
                        flash("File backup tidak valid. File archive.db tidak ditemukan.", "danger")
                        return redirect(url_for("restore_backup"))
                    backup_zip.extract("archive.db", tmp_dir)

                    upload_folder = os.path.realpath(app.config["UPLOAD_FOLDER"])
                    os.makedirs(upload_folder, exist_ok=True)
                    for info in backup_zip.infolist():
                        if info.is_dir():
                            continue
                        if info.filename.startswith("uploads/"):
                            relative_name = info.filename.split("uploads/", 1)[1]
                            target_path = os.path.join(upload_folder, *Path(relative_name).parts)
                            os.makedirs(os.path.dirname(target_path), exist_ok=True)
                            with backup_zip.open(info) as src, open(target_path, "wb") as dst:
                                shutil.copyfileobj(src, dst)

                db_ref = g.pop("db", None)
                if db_ref is not None:
                    db_ref.close()

                shutil.copy2(archive_path, app.config["DATABASE"])

            db = get_db()
            log_activity(db, "restore_backup", "Memulihkan database dari backup")
            db.commit()
            flash("Backup berhasil dipulihkan.", "success")
            return redirect(url_for("archives"))
        except zipfile.BadZipFile:
            flash("File backup yang dipilih bukan file ZIP yang valid.", "danger")
            return redirect(url_for("restore_backup"))
        except Exception as exc:
            app.logger.exception("Could not restore backup")
            flash(f"Restore backup gagal: {exc}", "danger")
            return redirect(url_for("restore_backup"))

    return render_template("backup_restore.html")


@app.route("/archives/trash")
def archive_trash():
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat mengelola tempat sampah arsip.", "danger")
        return redirect(url_for("dashboard"))

    records = get_db().execute(
        "SELECT * FROM archives WHERE deleted_at IS NOT NULL ORDER BY deleted_at DESC, id DESC"
    ).fetchall()
    return render_template("archive_trash.html", records=records)


@app.route("/archive/<int:archive_id>/restore", methods=["POST"])
def restore_archive(archive_id):
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat memulihkan arsip.", "danger")
        return redirect(url_for("dashboard"))

    db = get_db()
    row = db.execute(
        "SELECT nomor_surat, file_name FROM archives WHERE id = ? AND deleted_at IS NOT NULL",
        (archive_id,),
    ).fetchone()
    if row is None:
        flash("Arsip tidak ditemukan di tempat sampah.", "danger")
        return redirect(url_for("archive_trash"))

    db.execute("UPDATE archives SET deleted_at = NULL WHERE id = ?", (archive_id,))
    log_activity(
        db,
        "restore",
        f"Pulihkan arsip {row['nomor_surat'] or row['file_name'] or archive_id}",
        archive_id,
    )
    db.commit()
    flash("Arsip berhasil dipulihkan.", "success")
    return redirect(url_for("archive_trash"))


@app.route("/archive/<int:archive_id>/purge", methods=["POST"])
def purge_archive(archive_id):
    if session.get("role") != "admin":
        flash("Hanya admin yang dapat menghapus arsip permanen.", "danger")
        return redirect(url_for("dashboard"))

    db = get_db()
    row = db.execute(
        "SELECT file_path, file_name, nomor_surat FROM archives WHERE id = ? AND deleted_at IS NOT NULL",
        (archive_id,),
    ).fetchone()
    if row is None:
        flash("Arsip tidak ditemukan di tempat sampah.", "danger")
        return redirect(url_for("archive_trash"))

    db.execute("DELETE FROM archives WHERE id = ?", (archive_id,))
    remaining_references = db.execute(
        "SELECT COUNT(*) FROM archives WHERE file_path = ?", (row["file_path"],)
    ).fetchone()[0]
    log_activity(
        db,
        "purge",
        f"Hapus permanen arsip {row['nomor_surat'] or row['file_name'] or archive_id}",
        archive_id,
    )
    db.commit()

    file_path = os.path.realpath(row["file_path"] or "")
    upload_folder = os.path.realpath(app.config["UPLOAD_FOLDER"])
    try:
        is_uploaded_file = bool(row["file_path"]) and os.path.commonpath(
            (upload_folder, file_path)
        ) == upload_folder
    except ValueError:
        is_uploaded_file = False
    if is_uploaded_file and remaining_references == 0 and os.path.isfile(file_path):
        try:
            os.remove(file_path)
        except OSError:
            app.logger.exception("Could not remove uploaded file %s", file_path)

    flash("Arsip dihapus permanen.", "success")
    return redirect(url_for("archive_trash"))


@app.route("/archive/<int:archive_id>")
def archive_detail(archive_id):
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    row = db.execute(
        "SELECT * FROM archives WHERE id = ? AND deleted_at IS NULL", (archive_id,)
    ).fetchone()
    if row is None:
        flash("Dokumen tidak ditemukan.", "danger")
        return redirect(url_for("archives"))
    return render_template("archive_detail.html", record=row)


@app.route("/archive/<int:archive_id>/delete", methods=["POST"])
def archive_delete(archive_id):
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    row = db.execute(
        "SELECT file_path, file_name, nomor_surat FROM archives WHERE id = ? AND deleted_at IS NULL",
        (archive_id,),
    ).fetchone()
    if row is None:
        flash("Dokumen tidak ditemukan.", "danger")
        return redirect(url_for("archives"))

    db.execute(
        "UPDATE archives SET deleted_at = CURRENT_TIMESTAMP WHERE id = ?", (archive_id,)
    )
    log_activity(
        db,
        "trash",
        f"Pindahkan arsip {row['nomor_surat'] or row['file_name'] or archive_id} ke tempat sampah",
        archive_id,
    )
    db.commit()
    flash("Arsip dipindahkan ke tempat sampah dan dapat dipulihkan.", "success")
    return redirect(url_for("archives"))


@app.route("/archive/<int:archive_id>/ocr", methods=["POST"])
def archive_rerun_ocr(archive_id):
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    row = db.execute(
        "SELECT * FROM archives WHERE id = ? AND deleted_at IS NULL", (archive_id,)
    ).fetchone()
    if row is None:
        flash("Dokumen tidak ditemukan.", "danger")
        return redirect(url_for("archives"))

    if not os.path.isfile(row["file_path"]):
        flash("File dokumen tidak ditemukan di folder upload.", "danger")
        return redirect(url_for("archive_detail", archive_id=archive_id))

    with open(row["file_path"], "rb") as source:
        uploaded_file = FileStorage(stream=source, filename=row["file_name"])
        extracted_text = read_text_from_upload(uploaded_file)

    if extracted_text.startswith("OCR gambar gagal") or extracted_text.startswith("Dokumen berhasil diunggah"):
        flash(extracted_text, "warning")
        return redirect(url_for("archive_detail", archive_id=archive_id))

    metadata = extract_metadata(extracted_text)
    db.execute(
        """
        UPDATE archives
        SET nomor_surat = ?, tanggal = ?, perihal = ?, pengirim = ?, tujuan = ?, kategori = ?, extracted_text = ?
        WHERE id = ?
        """,
        (
            metadata["nomor_surat"] or row["nomor_surat"],
            metadata["tanggal"] or row["tanggal"],
            metadata["perihal"] or row["perihal"],
            metadata["pengirim"] or row["pengirim"],
            metadata["tujuan"] or row["tujuan"],
            classify_document(extracted_text),
            extracted_text,
            archive_id,
        ),
    )
    log_activity(db, "ocr_rerun", "OCR ulang arsip", archive_id)
    db.commit()
    flash("OCR ulang selesai. Periksa dan koreksi hasilnya bila diperlukan.", "success")
    return redirect(url_for("archive_detail", archive_id=archive_id))


@app.route("/archive/<int:archive_id>/edit", methods=["GET", "POST"])
def archive_edit(archive_id):
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    row = db.execute(
        "SELECT * FROM archives WHERE id = ? AND deleted_at IS NULL", (archive_id,)
    ).fetchone()
    if row is None:
        flash("Dokumen tidak ditemukan.", "danger")
        return redirect(url_for("archives"))

    if request.method == "POST":
        nomor_surat = request.form.get("nomor_surat", "").strip()
        tanggal = parse_date(request.form.get("tanggal", "").strip())
        perihal = request.form.get("perihal", "").strip()
        pengirim = request.form.get("pengirim", "").strip()
        tujuan = request.form.get("tujuan", "").strip()
        kategori = request.form.get("kategori", "").strip() or classify_document(request.form.get("extracted_text", ""))
        extracted_text = request.form.get("extracted_text", "").strip()

        db.execute(
            """
            UPDATE archives
            SET nomor_surat = ?, tanggal = ?, perihal = ?, pengirim = ?, tujuan = ?, kategori = ?, extracted_text = ?
            WHERE id = ?
            """,
            (nomor_surat, tanggal, perihal, pengirim, tujuan, kategori, extracted_text, archive_id),
        )
        log_activity(db, "edit", f"Koreksi data arsip {archive_id}", archive_id)
        db.commit()
        flash("Data arsip berhasil diperbarui.", "success")
        return redirect(url_for("archive_detail", archive_id=archive_id))

    return render_template("archive_edit.html", record=row)


@app.route("/reports")
def reports():
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    date_filters = get_date_filters()
    where_clause, params = archive_filter_query(
        start_date=date_filters["start_date"], end_date=date_filters["end_date"]
    )
    reports = db.execute(
        f"SELECT kategori, COUNT(*) AS total FROM archives {where_clause} GROUP BY kategori ORDER BY total DESC",
        params,
    ).fetchall()
    monthly_where = " AND ".join(
        [condition for condition in (where_clause.removeprefix("WHERE "), "tanggal IS NOT NULL", "tanggal != ''") if condition]
    )
    monthly = db.execute(
        f"SELECT substr(tanggal, 1, 7) as bulan, COUNT(*) as total FROM archives WHERE {monthly_where} GROUP BY substr(tanggal, 1, 7) ORDER BY bulan DESC",
        params,
    ).fetchall()
    return render_template("reports.html", reports=reports, monthly=monthly, **date_filters)


@app.route("/reports/export")
def export_reports():
    if "user" not in session:
        return redirect(url_for("login"))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    date_filters = get_date_filters()
    where_clause, params = archive_filter_query(
        start_date=date_filters["start_date"], end_date=date_filters["end_date"]
    )
    db = get_db()
    rows = db.execute(
        f"SELECT nomor_surat, tanggal, perihal, pengirim, tujuan, kategori, file_name FROM archives {where_clause} ORDER BY tanggal DESC, id DESC",
        params,
    ).fetchall()

    workbook = Workbook()
    detail_sheet = workbook.active
    detail_sheet.title = "Detail Arsip"
    headers = ("Nomor Surat", "Tanggal", "Perihal", "Pengirim", "Tujuan", "Kategori", "Nama File")
    detail_sheet.append(headers)
    for row in rows:
        detail_sheet.append(tuple(row[header] or "" for header in (
            "nomor_surat", "tanggal", "perihal", "pengirim", "tujuan", "kategori", "file_name"
        )))

    summary_sheet = workbook.create_sheet("Ringkasan")
    summary_sheet.append(("Kategori", "Jumlah Arsip"))
    summary_where, summary_params = archive_filter_query(
        start_date=date_filters["start_date"], end_date=date_filters["end_date"]
    )
    summary_rows = db.execute(
        f"SELECT kategori, COUNT(*) AS total FROM archives {summary_where} GROUP BY kategori ORDER BY total DESC",
        summary_params,
    ).fetchall()
    for row in summary_rows:
        summary_sheet.append((row["kategori"] or "Tanpa kategori", row["total"]))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="153B59")
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 14), 48)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="laporan_arsip.xlsx",
    )


@app.route("/download/<int:archive_id>")
def download_document(archive_id):
    if "user" not in session:
        return redirect(url_for("login"))

    db = get_db()
    row = db.execute(
        "SELECT * FROM archives WHERE id = ? AND deleted_at IS NULL", (archive_id,)
    ).fetchone()
    if row is None:
        flash("Dokumen tidak ditemukan.", "danger")
        return redirect(url_for("archives"))

    file_path = row["file_path"]
    return send_from_directory(os.path.dirname(file_path), os.path.basename(file_path), as_attachment=True)


@app.context_processor
def inject_user():
    return {"current_user": session.get("user"), "current_role": session.get("role")}


if __name__ == "__main__":
    debug_enabled = os.environ.get("FLASK_DEBUG", "0").lower() in ("1", "true", "yes")
    server_host = os.environ.get("FLASK_HOST", "127.0.0.1")
    app.run(debug=debug_enabled, host=server_host, port=5000)
